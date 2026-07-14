from __future__ import annotations

import base64
import binascii
import csv
import hashlib
import io
import json
import re
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import PurePath
from xml.etree import ElementTree

from openpyxl import load_workbook

from .models import (
    DocumentClause,
    DocumentSourceRecord,
    DocumentVersionCreateRequest,
)
from .risk_object_ingestion import configured_max_file_bytes


PARSER_VERSION = "document-parser-v2"
MAX_ARCHIVE_MEMBERS = 2_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_CLAUSES = 5_000

SUPPORTED_MEDIA_TYPES = {
    ".txt": {"text/plain"},
    ".md": {"text/markdown", "text/plain"},
    ".csv": {"text/csv", "application/csv", "text/plain"},
    ".xlsx": {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    },
    ".docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    },
    ".pdf": {"application/pdf"},
    ".png": {"image/png"},
    ".jpg": {"image/jpeg"},
    ".jpeg": {"image/jpeg"},
    ".tif": {"image/tiff"},
    ".tiff": {"image/tiff"},
}


class DocumentGovernanceError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedDocument:
    clauses: list[DocumentClause]
    parser_version: str
    parse_method: str
    result_hash: str


def build_document_source(
    version_id: str,
    request: DocumentVersionCreateRequest,
    created_at: datetime,
) -> DocumentSourceRecord:
    if request.file is None:
        assert request.content is not None
        payload = request.content.encode("utf-8")
        filename = f"{_safe_component(request.version_label)}.txt"
        media_type = "text/plain"
        source_hash = hashlib.sha256(payload).hexdigest()
    else:
        filename, suffix = _safe_filename(request.file.filename)
        media_type = request.file.media_type.split(";", 1)[0].strip().lower()
        allowed = SUPPORTED_MEDIA_TYPES.get(suffix)
        if not allowed or media_type not in allowed:
            raise DocumentGovernanceError(
                f"unsupported document media type {media_type!r} for {suffix or 'unknown'}"
            )
        try:
            payload = base64.b64decode(request.file.content_base64, validate=True)
        except (ValueError, binascii.Error) as exc:
            raise DocumentGovernanceError("document content is not valid base64") from exc
        if not payload:
            raise DocumentGovernanceError("document source file is empty")
        source_hash = hashlib.sha256(payload).hexdigest()
        if source_hash != request.file.sha256:
            raise DocumentGovernanceError(
                "document source SHA-256 does not match the uploaded content"
            )
    if len(payload) > configured_max_file_bytes():
        raise DocumentGovernanceError("document source file exceeds the size limit")
    return DocumentSourceRecord(
        version_id=version_id,
        filename=filename,
        media_type=media_type,
        source_hash=source_hash,
        source_size=len(payload),
        content_base64=base64.b64encode(payload).decode("ascii"),
        ocr_text=request.ocr_text,
        ocr_engine_version=request.ocr_engine_version,
        created_at=created_at,
    )


def parse_document_source(
    source: DocumentSourceRecord,
    *,
    parser_version: str = PARSER_VERSION,
) -> ParsedDocument:
    try:
        payload = base64.b64decode(source.content_base64, validate=True)
    except (ValueError, binascii.Error) as exc:
        raise DocumentGovernanceError("stored document source is not valid base64") from exc
    if hashlib.sha256(payload).hexdigest() != source.source_hash:
        raise DocumentGovernanceError("stored document source hash verification failed")
    suffix = PurePath(source.filename).suffix.lower()
    if suffix in {".txt", ".md"}:
        clauses = _parse_text(_decode_text(payload), source.filename)
        method = "plain_text"
    elif suffix == ".csv":
        clauses = _parse_csv(payload, source.filename)
        method = "csv_table"
    elif suffix == ".xlsx":
        _validate_zip(payload, "XLSX")
        clauses = _parse_xlsx(payload, source.filename)
        method = "xlsx_table"
    elif suffix == ".docx":
        _validate_zip(payload, "DOCX")
        clauses = _parse_docx(payload, source.filename)
        method = "docx_ooxml"
    elif suffix == ".pdf":
        clauses = _parse_pdf(payload, source)
        method = "pdf_text" if not source.ocr_text else "pdf_text_or_ocr"
    elif suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
        _validate_image(payload)
        clauses = _parse_ocr(source)
        method = "verified_ocr"
    else:
        raise DocumentGovernanceError(f"unsupported document format: {suffix}")
    clauses = clauses[:MAX_CLAUSES]
    if not clauses:
        raise DocumentGovernanceError("document parsing produced no indexable clauses")
    canonical = json.dumps(
        [item.model_dump(mode="json") for item in clauses],
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )
    return ParsedDocument(
        clauses=clauses,
        parser_version=parser_version,
        parse_method=method,
        result_hash=hashlib.sha256(canonical.encode("utf-8")).hexdigest(),
    )


def _safe_component(value: str) -> str:
    normalized = re.sub(r"[^A-Za-z0-9._-]+", "-", value.strip()).strip(".-")
    return normalized[:80] or "document"


def _safe_filename(filename: str) -> tuple[str, str]:
    normalized = filename.strip()
    if not normalized or PurePath(normalized).name != normalized:
        raise DocumentGovernanceError("document filename must not contain a path")
    if "\x00" in normalized:
        raise DocumentGovernanceError("document filename contains a null byte")
    suffix = PurePath(normalized).suffix.lower()
    if suffix not in SUPPORTED_MEDIA_TYPES:
        raise DocumentGovernanceError(
            "supported document formats are TXT, Markdown, CSV, XLSX, DOCX, PDF and scanned images"
        )
    return normalized, suffix


def _validate_zip(payload: bytes, label: str) -> None:
    if not payload.startswith(b"PK\x03\x04"):
        raise DocumentGovernanceError(f"{label} input is not an OOXML ZIP package")
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS:
                raise DocumentGovernanceError(
                    f"{label} archive contains too many members"
                )
            total = 0
            for member in members:
                path = PurePath(member.filename)
                if path.is_absolute() or ".." in path.parts:
                    raise DocumentGovernanceError(
                        f"{label} archive contains an unsafe member path"
                    )
                total += member.file_size
                if total > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                    raise DocumentGovernanceError(
                        f"{label} archive expands beyond the safety limit"
                    )
    except zipfile.BadZipFile as exc:
        raise DocumentGovernanceError(f"{label} archive is invalid") from exc


def _validate_image(payload: bytes) -> None:
    try:
        from PIL import Image

        with Image.open(io.BytesIO(payload)) as image:
            image.verify()
    except Exception as exc:
        raise DocumentGovernanceError("scanned document image is invalid") from exc


def _decode_text(payload: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentGovernanceError("text document is not UTF-8 or GB18030")


def _make_clause(
    index: int,
    text: str,
    *,
    heading: str,
    source_locator: str,
    page_number: int | None = None,
    section_path: list[str] | None = None,
    table_name: str | None = None,
    row_start: int | None = None,
    row_end: int | None = None,
    content_type: str = "paragraph",
) -> DocumentClause:
    normalized = re.sub(r"\s+", " ", text).strip()
    clause_match = re.match(r"^(第[^，。；:：]{1,20}条|[0-9]+(?:\.[0-9]+)*)", normalized)
    return DocumentClause(
        clause_id=f"C{index:04d}",
        heading=heading[:100] or f"条款 {index}",
        text=normalized,
        page_number=page_number,
        section_path=section_path or [],
        clause_number=clause_match.group(1) if clause_match else None,
        table_name=table_name,
        row_start=row_start,
        row_end=row_end,
        source_locator=source_locator,
        content_type=content_type,
        text_hash=hashlib.sha256(normalized.encode("utf-8")).hexdigest(),
    )


def _paragraphs(text: str) -> list[str]:
    return [
        item.strip()
        for item in re.split(r"\n\s*\n|(?<=。)\s*\n", text.replace("\r\n", "\n"))
        if item.strip()
    ]


def _parse_text(text: str, filename: str) -> list[DocumentClause]:
    clauses: list[DocumentClause] = []
    for page_number, page_text in enumerate(text.split("\f"), start=1):
        section_path: list[str] = []
        for paragraph in _paragraphs(page_text):
            first_line = paragraph.splitlines()[0].strip()
            if re.match(r"^(第.+[章节]|[0-9]+(?:\.[0-9]+)*\s+)", first_line):
                section_path = [first_line[:100]]
            clauses.append(
                _make_clause(
                    len(clauses) + 1,
                    paragraph,
                    heading=first_line,
                    source_locator=f"{filename}#page={page_number}&clause={len(clauses) + 1}",
                    page_number=page_number,
                    section_path=section_path,
                )
            )
    return clauses


def _parse_csv(payload: bytes, filename: str) -> list[DocumentClause]:
    rows = list(csv.reader(io.StringIO(_decode_text(payload))))
    return _table_rows_to_clauses(rows, filename, "CSV")


def _table_rows_to_clauses(
    rows: list[list[object]], filename: str, table_name: str
) -> list[DocumentClause]:
    clauses: list[DocumentClause] = []
    for row_number, row in enumerate(rows, start=1):
        values = [str(value).strip() if value is not None else "" for value in row]
        if not any(values):
            continue
        text = " | ".join(values)
        clauses.append(
            _make_clause(
                len(clauses) + 1,
                text,
                heading=f"{table_name} 第 {row_number} 行",
                source_locator=f"{filename}#{table_name}!R{row_number}",
                table_name=table_name,
                row_start=row_number,
                row_end=row_number,
                content_type="table_row",
            )
        )
    return clauses


def _parse_xlsx(payload: bytes, filename: str) -> list[DocumentClause]:
    try:
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:
        raise DocumentGovernanceError("XLSX workbook cannot be parsed") from exc
    clauses: list[DocumentClause] = []
    try:
        for worksheet in workbook.worksheets:
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            for clause in _table_rows_to_clauses(rows, filename, worksheet.title):
                clauses.append(
                    clause.model_copy(
                        update={"clause_id": f"C{len(clauses) + 1:04d}"}
                    )
                )
    finally:
        workbook.close()
    return clauses


def _docx_text(element: ElementTree.Element) -> str:
    parts = []
    for child in element.iter():
        name = child.tag.rsplit("}", 1)[-1]
        if name == "t" and child.text:
            parts.append(child.text)
        elif name == "tab":
            parts.append("\t")
    return "".join(parts).strip()


def _parse_docx(payload: bytes, filename: str) -> list[DocumentClause]:
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            xml_payload = archive.read("word/document.xml")
    except (KeyError, zipfile.BadZipFile) as exc:
        raise DocumentGovernanceError("DOCX is missing word/document.xml") from exc
    try:
        root = ElementTree.fromstring(xml_payload)
    except ElementTree.ParseError as exc:
        raise DocumentGovernanceError("DOCX document XML is invalid") from exc
    body = next((item for item in root.iter() if item.tag.endswith("}body")), None)
    if body is None:
        raise DocumentGovernanceError("DOCX body is missing")
    clauses: list[DocumentClause] = []
    section_path: list[str] = []
    page_number = 1
    table_number = 0
    for child in body:
        name = child.tag.rsplit("}", 1)[-1]
        if name == "p":
            text = _docx_text(child)
            if text:
                if re.match(r"^(第.+[章节]|[0-9]+(?:\.[0-9]+)*\s+)", text):
                    section_path = [text[:100]]
                clauses.append(
                    _make_clause(
                        len(clauses) + 1,
                        text,
                        heading=text[:100],
                        source_locator=f"{filename}#page={page_number}&clause={len(clauses) + 1}",
                        page_number=page_number,
                        section_path=section_path,
                    )
                )
            page_number += sum(
                1
                for item in child.iter()
                if item.tag.endswith("}br")
                and item.attrib.get(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type"
                )
                == "page"
            )
        elif name == "tbl":
            table_number += 1
            row_number = 0
            for row in (item for item in child if item.tag.endswith("}tr")):
                row_number += 1
                values = [
                    _docx_text(cell)
                    for cell in row
                    if cell.tag.endswith("}tc")
                ]
                if not any(values):
                    continue
                clauses.append(
                    _make_clause(
                        len(clauses) + 1,
                        " | ".join(values),
                        heading=f"表 {table_number} 第 {row_number} 行",
                        source_locator=f"{filename}#table={table_number}&row={row_number}",
                        page_number=page_number,
                        section_path=section_path,
                        table_name=f"Table {table_number}",
                        row_start=row_number,
                        row_end=row_number,
                        content_type="table_row",
                    )
                )
    return clauses


def _parse_pdf(payload: bytes, source: DocumentSourceRecord) -> list[DocumentClause]:
    if not payload.startswith(b"%PDF-"):
        raise DocumentGovernanceError("PDF input does not have a valid PDF signature")
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        if source.ocr_text:
            return _parse_ocr(source)
        raise DocumentGovernanceError(
            "PDF parser dependency is unavailable and no verified OCR text was supplied"
        ) from exc
    try:
        reader = PdfReader(io.BytesIO(payload), strict=False)
        if reader.is_encrypted and not reader.decrypt(""):
            raise DocumentGovernanceError("encrypted PDF documents are not supported")
        clauses: list[DocumentClause] = []
        for page_number, page in enumerate(reader.pages, start=1):
            for paragraph in _paragraphs(page.extract_text() or ""):
                clauses.append(
                    _make_clause(
                        len(clauses) + 1,
                        paragraph,
                        heading=paragraph[:100],
                        source_locator=f"{source.filename}#page={page_number}&clause={len(clauses) + 1}",
                        page_number=page_number,
                    )
                )
    except DocumentGovernanceError:
        raise
    except Exception as exc:
        if source.ocr_text:
            return _parse_ocr(source)
        raise DocumentGovernanceError("PDF document cannot be parsed") from exc
    if clauses:
        return clauses
    if source.ocr_text:
        return _parse_ocr(source)
    raise DocumentGovernanceError(
        "scanned PDF contains no extractable text; supply verified ocr_text and ocr_engine_version"
    )


def _parse_ocr(source: DocumentSourceRecord) -> list[DocumentClause]:
    if not source.ocr_text or not source.ocr_engine_version:
        raise DocumentGovernanceError(
            "scanned documents require verified ocr_text and ocr_engine_version"
        )
    clauses = _parse_text(source.ocr_text, source.filename)
    return [
        item.model_copy(
            update={
                "content_type": "ocr_paragraph",
                "source_locator": (
                    f"{item.source_locator}&ocr_engine={source.ocr_engine_version}"
                ),
            }
        )
        for item in clauses
    ]
