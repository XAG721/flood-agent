from __future__ import annotations

import base64
import binascii
import csv
import hashlib
import io
import json
import os
import re
import zipfile
from dataclasses import dataclass
from pathlib import PurePosixPath
from typing import Any

from openpyxl import load_workbook
from pydantic import ValidationError

from .models import (
    IngestionFileEnvelope,
    RiskObjectImportQuarantineItem,
    RiskObjectInput,
)


DEFAULT_MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 2_000
MAX_COMPRESSION_RATIO = 100
MAX_ROWS = 5_000
MAX_COLUMNS = 100
MAX_CELL_CHARACTERS = 20_000
MAX_JSON_NESTING_DEPTH = 64

SUPPORTED_MEDIA_TYPES = {
    ".csv": {"text/csv", "application/csv", "application/octet-stream"},
    ".xlsx": {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/octet-stream",
    },
    ".json": {"application/json", "text/json", "application/octet-stream"},
    ".geojson": {
        "application/geo+json",
        "application/json",
        "application/octet-stream",
    },
}

COLUMN_ALIASES = {
    "对象编号": "object_id",
    "对象id": "object_id",
    "规范对象编号": "canonical_object_id",
    "对象名称": "name",
    "对象类型": "object_type",
    "位置": "location",
    "经度": "longitude",
    "纬度": "latitude",
    "责任单位": "responsible_organization",
    "责任岗位": "responsible_role",
    "触发原因": "trigger_reasons",
    "来源引用": "source_refs",
    "脆弱性": "vulnerability",
    "历史风险": "historical_risk",
    "风险分": "risk_score",
    "风险评分": "risk_score",
    "系统说明": "system_explanation",
    "别名": "aliases",
    "台账状态": "registry_status",
    "有效开始": "registry_valid_from",
    "有效截止": "registry_valid_until",
    "模拟数据": "is_simulated",
}

LIST_FIELDS = {
    "aliases",
    "trigger_reasons",
    "source_refs",
    "missing_fields",
}

REQUIRED_FIELDS = {
    "object_id",
    "name",
    "object_type",
    "responsible_organization",
    "responsible_role",
    "risk_score",
}


class RiskObjectIngestionError(ValueError):
    """Raised when a source file fails the security or format boundary."""


@dataclass(frozen=True)
class ParsedRiskObjectFile:
    source_format: str
    source_hash: str
    source_bytes: int
    objects: list[RiskObjectInput]
    quarantine: list[RiskObjectImportQuarantineItem]


def configured_max_file_bytes() -> int:
    try:
        configured = int(
            os.getenv("FLOOD_INGESTION_MAX_FILE_BYTES", str(DEFAULT_MAX_FILE_BYTES))
        )
    except ValueError:
        return DEFAULT_MAX_FILE_BYTES
    return min(50 * 1024 * 1024, max(1024, configured))


def _safe_filename(filename: str) -> tuple[str, str]:
    normalized = filename.strip()
    if (
        not normalized
        or normalized in {".", ".."}
        or "/" in normalized
        or "\\" in normalized
        or any(ord(character) < 32 for character in normalized)
    ):
        raise RiskObjectIngestionError(
            "source filename must be a plain filename without path components"
        )
    suffix = PurePosixPath(normalized).suffix.lower()
    if suffix not in SUPPORTED_MEDIA_TYPES:
        raise RiskObjectIngestionError(
            "risk-object imports support only CSV, XLSX, JSON and GeoJSON files"
        )
    return normalized, suffix


def decode_file_envelope(envelope: IngestionFileEnvelope) -> tuple[bytes, str]:
    _, suffix = _safe_filename(envelope.filename)
    media_type = envelope.media_type.split(";", 1)[0].strip().lower()
    if media_type not in SUPPORTED_MEDIA_TYPES[suffix]:
        raise RiskObjectIngestionError(
            f"media type {media_type!r} does not match {suffix} risk-object input"
        )
    try:
        payload = base64.b64decode(envelope.content_base64, validate=True)
    except (ValueError, binascii.Error) as exc:
        raise RiskObjectIngestionError("file content is not valid base64") from exc
    if not payload:
        raise RiskObjectIngestionError("risk-object import file is empty")
    if len(payload) > configured_max_file_bytes():
        raise RiskObjectIngestionError("risk-object import file exceeds the size limit")
    actual_hash = hashlib.sha256(payload).hexdigest()
    if actual_hash != envelope.sha256:
        raise RiskObjectIngestionError(
            "risk-object import SHA-256 does not match the uploaded content"
        )
    return payload, suffix


def _validate_xlsx_archive(payload: bytes) -> None:
    if not payload.startswith(b"PK\x03\x04"):
        raise RiskObjectIngestionError("XLSX input is not an OOXML ZIP package")
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS:
                raise RiskObjectIngestionError("XLSX archive contains too many members")
            total_uncompressed = 0
            for member in members:
                path = PurePosixPath(member.filename)
                if path.is_absolute() or ".." in path.parts:
                    raise RiskObjectIngestionError(
                        "XLSX archive contains an unsafe member path"
                    )
                total_uncompressed += member.file_size
                if total_uncompressed > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                    raise RiskObjectIngestionError(
                        "XLSX archive exceeds the uncompressed size limit"
                    )
                if (
                    member.file_size > 1024
                    and member.compress_size > 0
                    and member.file_size / member.compress_size > MAX_COMPRESSION_RATIO
                ):
                    raise RiskObjectIngestionError(
                        "XLSX archive exceeds the compression ratio limit"
                    )
                lowered = member.filename.lower()
                if lowered.endswith("vbaproject.bin"):
                    raise RiskObjectIngestionError(
                        "macro-enabled workbooks are forbidden"
                    )
                if lowered.startswith("xl/externallinks/"):
                    raise RiskObjectIngestionError(
                        "workbooks with external links are forbidden"
                    )
                if lowered.startswith("xl/worksheets/") and lowered.endswith(".xml"):
                    worksheet_xml = archive.read(member)
                    if b"<f" in worksheet_xml:
                        raise RiskObjectIngestionError(
                            "workbook formulas are forbidden in risk-object imports"
                        )
            names = {member.filename for member in members}
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise RiskObjectIngestionError(
                    "XLSX archive is missing required workbook members"
                )
    except (zipfile.BadZipFile, OSError) as exc:
        raise RiskObjectIngestionError("XLSX input is not a valid ZIP package") from exc


def _decode_text(payload: bytes, *, source_format: str) -> str:
    if b"\x00" in payload:
        raise RiskObjectIngestionError(
            f"{source_format} input contains binary NUL characters"
        )
    try:
        return payload.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise RiskObjectIngestionError(
            f"{source_format} input must use UTF-8 encoding"
        ) from exc


def _normalize_header(value: Any) -> str:
    header = str(value or "").strip().lstrip("\ufeff")
    lowered = header.lower()
    return COLUMN_ALIASES.get(lowered, COLUMN_ALIASES.get(header, lowered))


def _csv_rows(payload: bytes) -> list[tuple[int, dict[str, Any]]]:
    text = _decode_text(payload, source_format="CSV")
    try:
        reader = csv.DictReader(io.StringIO(text), strict=True)
        if reader.fieldnames is None:
            raise RiskObjectIngestionError("CSV input is missing a header row")
        if len(reader.fieldnames) > MAX_COLUMNS:
            raise RiskObjectIngestionError("CSV input contains too many columns")
        normalized_headers = [_normalize_header(item) for item in reader.fieldnames]
        if not all(normalized_headers) or len(set(normalized_headers)) != len(
            normalized_headers
        ):
            raise RiskObjectIngestionError("CSV headers must be non-empty and unique")
        rows: list[tuple[int, dict[str, Any]]] = []
        for row_number, raw in enumerate(reader, start=2):
            if row_number > MAX_ROWS + 1:
                raise RiskObjectIngestionError("CSV input exceeds the row limit")
            row = {
                normalized_headers[index]: raw.get(original)
                for index, original in enumerate(reader.fieldnames)
            }
            if any(str(value or "").strip() for value in row.values()):
                rows.append((row_number, row))
        return rows
    except csv.Error as exc:
        raise RiskObjectIngestionError("CSV input is malformed") from exc


def _xlsx_rows(payload: bytes) -> list[tuple[int, dict[str, Any]]]:
    _validate_xlsx_archive(payload)
    try:
        workbook = load_workbook(
            io.BytesIO(payload),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception as exc:
        raise RiskObjectIngestionError("XLSX workbook cannot be parsed safely") from exc
    try:
        if not workbook.worksheets:
            raise RiskObjectIngestionError("XLSX workbook has no worksheets")
        worksheet = workbook.worksheets[0]
        iterator = worksheet.iter_rows(values_only=True)
        try:
            headers_raw = next(iterator)
        except StopIteration as exc:
            raise RiskObjectIngestionError("XLSX input is empty") from exc
        if len(headers_raw) > MAX_COLUMNS:
            raise RiskObjectIngestionError("XLSX input contains too many columns")
        headers = [_normalize_header(item) for item in headers_raw]
        if not all(headers) or len(set(headers)) != len(headers):
            raise RiskObjectIngestionError("XLSX headers must be non-empty and unique")
        rows: list[tuple[int, dict[str, Any]]] = []
        for row_number, values in enumerate(iterator, start=2):
            if row_number > MAX_ROWS + 1:
                raise RiskObjectIngestionError("XLSX input exceeds the row limit")
            row = {
                header: values[index] if index < len(values) else None
                for index, header in enumerate(headers)
            }
            if any(str(value or "").strip() for value in row.values()):
                rows.append((row_number, row))
        return rows
    finally:
        workbook.close()


def _json_rows(payload: bytes, *, geojson: bool) -> list[tuple[int, dict[str, Any]]]:
    text = _decode_text(payload, source_format="GeoJSON" if geojson else "JSON")
    try:
        document = json.loads(text)
    except (json.JSONDecodeError, RecursionError) as exc:
        raise RiskObjectIngestionError("JSON input is malformed") from exc
    stack = [(document, 1)]
    while stack:
        item, depth = stack.pop()
        if depth > MAX_JSON_NESTING_DEPTH:
            raise RiskObjectIngestionError("JSON input exceeds the nesting depth limit")
        if isinstance(item, dict):
            stack.extend((value, depth + 1) for value in item.values())
        elif isinstance(item, list):
            stack.extend((value, depth + 1) for value in item)
    if geojson:
        if (
            not isinstance(document, dict)
            or document.get("type") != "FeatureCollection"
        ):
            raise RiskObjectIngestionError("GeoJSON input must be a FeatureCollection")
        features = document.get("features")
        if not isinstance(features, list):
            raise RiskObjectIngestionError("GeoJSON features must be an array")
        if len(features) > MAX_ROWS:
            raise RiskObjectIngestionError("GeoJSON input exceeds the feature limit")
        rows: list[tuple[int, dict[str, Any]]] = []
        for index, feature in enumerate(features, start=1):
            if not isinstance(feature, dict):
                rows.append((index, {"_row_error": "feature must be an object"}))
                continue
            properties = feature.get("properties")
            geometry = feature.get("geometry")
            if not isinstance(properties, dict):
                rows.append(
                    (index, {"_row_error": "feature properties must be an object"})
                )
                continue
            row = dict(properties)
            if not isinstance(geometry, dict) or geometry.get("type") != "Point":
                row["_row_error"] = "risk-object GeoJSON geometry must be Point"
            else:
                coordinates = geometry.get("coordinates")
                if not isinstance(coordinates, list) or len(coordinates) < 2:
                    row["_row_error"] = "GeoJSON Point coordinates are invalid"
                else:
                    row["longitude"] = coordinates[0]
                    row["latitude"] = coordinates[1]
            rows.append((index, row))
        return rows

    if isinstance(document, dict):
        document = document.get("objects")
    if not isinstance(document, list):
        raise RiskObjectIngestionError(
            "JSON input must be an array or an object containing an objects array"
        )
    if len(document) > MAX_ROWS:
        raise RiskObjectIngestionError("JSON input exceeds the object limit")
    return [
        (
            index,
            item
            if isinstance(item, dict)
            else {"_row_error": "item must be an object"},
        )
        for index, item in enumerate(document, start=1)
    ]


def _split_list(value: Any) -> list[str]:
    if value is None or value == "":
        return []
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    return [
        item.strip() for item in re.split(r"[|；;、\n]+", str(value)) if item.strip()
    ]


def _parse_boolean(value: Any, *, default: bool = True) -> bool:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().lower()
    if normalized in {"1", "true", "yes", "y", "是", "模拟"}:
        return True
    if normalized in {"0", "false", "no", "n", "否", "真实"}:
        return False
    raise ValueError("is_simulated must be a boolean value")


def _coerce_row(
    raw: dict[str, Any],
    *,
    row_number: int,
    source_hash: str,
    source_version: str,
) -> RiskObjectInput:
    if raw.get("_row_error"):
        raise ValueError(str(raw["_row_error"]))
    row = {_normalize_header(key): value for key, value in raw.items()}
    for key, value in row.items():
        if isinstance(value, str) and len(value) > MAX_CELL_CHARACTERS:
            raise ValueError(f"field {key} exceeds the cell length limit")
    missing_required = sorted(
        field for field in REQUIRED_FIELDS if not str(row.get(field, "")).strip()
    )
    if missing_required:
        raise ValueError("missing required fields: " + ", ".join(missing_required))
    for field in LIST_FIELDS:
        if field in row:
            row[field] = _split_list(row[field])
    longitude = row.get("longitude")
    latitude = row.get("latitude")
    if longitude not in (None, ""):
        row["longitude"] = float(longitude)
    else:
        row["longitude"] = None
    if latitude not in (None, ""):
        row["latitude"] = float(latitude)
    else:
        row["latitude"] = None
    row["risk_score"] = float(row["risk_score"])
    row["is_simulated"] = _parse_boolean(row.get("is_simulated"), default=True)
    row.setdefault("canonical_object_id", None)
    row.setdefault("aliases", [])
    row.setdefault("duplicate_of", None)
    row.setdefault(
        "location",
        (
            f"EPSG:4326({row['longitude']}, {row['latitude']})"
            if row["longitude"] is not None
            else "未提供文字位置"
        ),
    )
    row.setdefault("trigger_reasons", ["风险对象台账导入，待候选算法与人工核验"])
    row.setdefault("source_refs", [f"registry-file:{source_hash}:row:{row_number}"])
    row.setdefault("vulnerability", "台账未登记额外脆弱性")
    row.setdefault("historical_risk", "")
    row.setdefault(
        "system_explanation",
        "风险对象来自受控文件导入；候选算法结果仍需人工确认。",
    )
    row.setdefault("source_type", "file_import")
    row["source_version"] = str(row.get("source_version") or source_version)
    row["data_version"] = str(
        row.get("data_version") or f"{source_version}:{source_hash[:12]}"
    )
    row.setdefault("missing_fields", [])
    row.setdefault("calibration_version", "candidate-logistic-v1")
    row.setdefault("association_mode", "registry_import")
    row.setdefault("registry_status", "active")
    allowed = set(RiskObjectInput.model_fields)
    return RiskObjectInput.model_validate(
        {key: value for key, value in row.items() if key in allowed}
    )


def parse_risk_object_file(
    envelope: IngestionFileEnvelope,
    *,
    source_version: str,
) -> ParsedRiskObjectFile:
    payload, suffix = decode_file_envelope(envelope)
    source_hash = hashlib.sha256(payload).hexdigest()
    if suffix == ".csv":
        rows = _csv_rows(payload)
        source_format = "csv"
    elif suffix == ".xlsx":
        rows = _xlsx_rows(payload)
        source_format = "xlsx"
    elif suffix == ".geojson":
        rows = _json_rows(payload, geojson=True)
        source_format = "geojson"
    else:
        rows = _json_rows(payload, geojson=False)
        source_format = "json"

    objects: list[RiskObjectInput] = []
    quarantine: list[RiskObjectImportQuarantineItem] = []
    object_ids: set[str] = set()
    for row_number, raw in rows:
        source_id = str(raw.get("object_id") or raw.get("对象编号") or "").strip()
        try:
            item = _coerce_row(
                raw,
                row_number=row_number,
                source_hash=source_hash,
                source_version=source_version,
            )
            if item.object_id in object_ids:
                raise ValueError(f"duplicate object_id in one import: {item.object_id}")
            object_ids.add(item.object_id)
            objects.append(item)
        except (TypeError, ValueError, ValidationError) as exc:
            quarantine.append(
                RiskObjectImportQuarantineItem(
                    source_row=row_number,
                    source_id=source_id,
                    reason_code=(
                        "DUPLICATE_OBJECT_ID"
                        if "duplicate object_id" in str(exc)
                        else "INVALID_RISK_OBJECT"
                    ),
                    reason=str(exc),
                )
            )
    return ParsedRiskObjectFile(
        source_format=source_format,
        source_hash=source_hash,
        source_bytes=len(payload),
        objects=objects,
        quarantine=quarantine,
    )
